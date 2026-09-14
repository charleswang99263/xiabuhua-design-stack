#!/usr/bin/env python3
"""Create a basic, position-aware inventory from a .pptx file.

This intentionally reports a narrow inventory rather than a complete
PowerPoint conversion. It reports slide/shape order, text where python-pptx can
read it, geometry, selected image metadata, and unsupported shape categories.
The legacy binary .ppt format is rejected.

Usage:
    python extract-pptx.py input.pptx [output_dir]

Requires an already available ``python-pptx`` installation. This script does
not install dependencies.
"""

import json
import os
import sys
from pathlib import Path


SUPPORTED_PPTX = ".pptx"


def _geometry(shape):
    return {
        "left_emu": int(shape.left),
        "top_emu": int(shape.top),
        "width_emu": int(shape.width),
        "height_emu": int(shape.height),
        "left_in": round(shape.left / 914400, 4),
        "top_in": round(shape.top / 914400, 4),
        "width_in": round(shape.width / 914400, 4),
        "height_in": round(shape.height / 914400, 4),
    }


def _shape_type_name(shape):
    return getattr(shape.shape_type, "name", str(shape.shape_type))


def extract_pptx(file_path, output_dir="."):
    """Return a basic inventory and write extracted images when available."""
    source = Path(file_path)
    if source.suffix.lower() != SUPPORTED_PPTX:
        raise ValueError("Only .pptx is supported; legacy .ppt files are rejected.")

    try:
        from pptx import Presentation
        from pptx.enum.shapes import MSO_SHAPE_TYPE
    except ImportError as error:
        raise RuntimeError("python-pptx is required and must already be installed") from error

    prs = Presentation(str(source))
    destination = Path(output_dir)
    assets_dir = destination / "assets"
    assets_dir.mkdir(parents=True, exist_ok=True)
    unsupported = set()
    slides = []
    picture_type = MSO_SHAPE_TYPE.PICTURE

    for slide_number, slide in enumerate(prs.slides, start=1):
        shapes = []
        for order, shape in enumerate(slide.shapes, start=1):
            kind = _shape_type_name(shape)
            item = {
                "order": order,
                "shape_id": int(shape.shape_id),
                "name": shape.name,
                "type": kind,
                **_geometry(shape),
            }
            if getattr(shape, "has_text_frame", False):
                item["text"] = shape.text

            if shape.shape_type == picture_type:
                image = shape.image
                image_name = f"slide{slide_number}_img{len([s for s in shapes if s.get('type') == 'PICTURE']) + 1}.{image.ext}"
                image_path = assets_dir / image_name
                image_path.write_bytes(image.blob)
                item["asset"] = str(Path("assets") / image_name)
            else:
                # Tables, charts, groups, SmartArt, and other advanced shapes
                # are inventoried by geometry but are not flattened here.
                supported_text = getattr(shape, "has_text_frame", False)
                if not supported_text:
                    unsupported.add(kind)
                    item["unsupported"] = kind
            shapes.append(item)

        slides.append({
            "number": slide_number,
            "order": slide_number,
            "shape_count": len(shapes),
            "shapes": shapes,
        })

    return {
        "inventory_version": "1.0",
        "scope": "basic .pptx slide and shape inventory",
        "source": source.name,
        "format": "pptx",
        "slide_size": {
            "width_emu": int(prs.slide_width),
            "height_emu": int(prs.slide_height),
            "width_in": round(prs.slide_width / 914400, 4),
            "height_in": round(prs.slide_height / 914400, 4),
        },
        "slide_count": len(slides),
        "slides": slides,
        "unsupported_categories": sorted(unsupported),
    }


def main(argv):
    if len(argv) < 2 or len(argv) > 3:
        print("Usage: python extract-pptx.py <input.pptx> [output_dir]", file=sys.stderr)
        return 2
    try:
        inventory = extract_pptx(argv[1], argv[2] if len(argv) == 3 else ".")
    except (OSError, RuntimeError, ValueError) as error:
        print(f"ERROR: {error}", file=sys.stderr)
        return 2

    output_dir = Path(argv[2] if len(argv) == 3 else ".")
    output_dir.mkdir(parents=True, exist_ok=True)
    output_path = output_dir / "extracted-slides.json"
    output_path.write_text(json.dumps(inventory, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    print(f"Wrote basic PPTX inventory for {inventory['slide_count']} slides to {output_path}")
    if inventory["unsupported_categories"]:
        print("Unsupported categories: " + ", ".join(inventory["unsupported_categories"]))
    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv))
